#!/usr/bin/env python3
"""Shared table renderer for hx-pptx slides.

Renders styled data tables onto PowerPoint slides with hx theme colours,
dynamic font sizing, clean borders, and <br> multi-line cell support.

Used by generate.py (via the <!-- table: path.json --> directive) and by
domain-specific skills like create-pain-slide that delegate rendering here.
"""

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Cm, Pt

# ── hx theme colours ─────────────────────────────────────────────────────────

DARK_TEXT = RGBColor(0x1C, 0x27, 0x33)     # dk1
LIGHT_BG = RGBColor(0xF5, 0xF5, 0xF5)     # lt2
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SEPARATOR = RGBColor(0xD0, 0xD0, 0xD0)

_HEADER_COLORS = {
    "blue": RGBColor(0x4D, 0x6F, 0xF8),       # accent2
    "orange": RGBColor(0xEE, 0x6C, 0x5A),     # accent4
    "green": RGBColor(0x01, 0x51, 0x4F),       # dk2
}

# ── Default positioning (full-width, below headline area) ────────────────────

DEFAULT_POSITION = {
    "left_cm": 1.26,
    "top_cm": 5.5,
    "width_cm": 31.35,
    "height_cm": 12.5,
}

# ── Dynamic font sizing by data-row count ────────────────────────────────────

_FONT_SIZES = {
    1: {"header": 12, "body": 10},
    2: {"header": 12, "body": 10},
    3: {"header": 11, "body": 9},
    4: {"header": 11, "body": 8},
    5: {"header": 10, "body": 7.5},
    6: {"header": 10, "body": 7},
    7: {"header": 9,  "body": 7},
}


def _get_font_sizes(row_count):
    clamped = max(1, min(row_count, max(_FONT_SIZES)))
    return _FONT_SIZES[clamped]


# ── OOXML border helpers ─────────────────────────────────────────────────────

_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _remove_cell_borders(cell):
    """Strip all borders from a table cell."""
    tc_pr = cell._tc.find(f'{{{_NS}}}tcPr')
    if tc_pr is None:
        tc_pr = etree.SubElement(cell._tc, f'{{{_NS}}}tcPr')

    for side in ('lnL', 'lnR', 'lnT', 'lnB', 'lnTlToBr', 'lnBlToTr'):
        for existing in tc_pr.findall(f'{{{_NS}}}{side}'):
            tc_pr.remove(existing)

    # Insert no-fill borders at the front (OOXML sequence order).
    for i, side in enumerate(('lnL', 'lnR', 'lnT', 'lnB')):
        ln = etree.Element(f'{{{_NS}}}{side}')
        ln.set('w', '12700')
        etree.SubElement(ln, f'{{{_NS}}}noFill')
        tc_pr.insert(i, ln)


def _set_thin_border(cell, sides=None, color=SEPARATOR):
    """Add thin (0.5 pt) solid borders on the specified sides."""
    if sides is None:
        sides = ["lnB"]
    tcPr = cell._tc.get_or_add_tcPr()
    for side in sides:
        border = tcPr.find(f'{{{_NS}}}{side}')
        if border is None:
            border = tcPr.makeelement(f'{{{_NS}}}{side}', {})
            tcPr.append(border)
        border.set('w', '6350')   # 0.5 pt
        border.set('cap', 'flat')
        border.set('cmpd', 'sng')
        for child in list(border):
            border.remove(child)
        solidFill = border.makeelement(f'{{{_NS}}}solidFill', {})
        srgbClr = solidFill.makeelement(
            f'{{{_NS}}}srgbClr', {'val': str(color)},
        )
        solidFill.append(srgbClr)
        border.append(solidFill)


# ── Cell styling ─────────────────────────────────────────────────────────────

def _style_cell(cell, text, font_size, *, bold=False, bg_color=None,
                font_color=DARK_TEXT, font_name="Arial", v_anchor=MSO_ANCHOR.TOP):
    """Populate and style a single cell.  Supports ``<br>`` for line breaks."""
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.15)
    tf.margin_right = Cm(0.15)
    tf.margin_top = Cm(0.08)
    tf.margin_bottom = Cm(0.08)
    cell.vertical_anchor = v_anchor

    parts = text.split("<br>") if "<br>" in text else [text]
    for i, part in enumerate(parts):
        part = part.strip()
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = part
        p.font.size = Pt(font_size)
        p.font.name = font_name
        p.font.color.rgb = font_color
        p.font.bold = bold
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        p.alignment = PP_ALIGN.LEFT

    if bg_color:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg_color


# ── Public API ───────────────────────────────────────────────────────────────

def render_table(slide, table_def, colour=None, position=None):
    """Render a styled data table onto *slide*.

    Parameters
    ----------
    slide : pptx.slide.Slide
        The slide to draw on.
    table_def : dict
        Accepted keys:

        ``columns``
            ``[{"header": str, "width": float}, ...]`` — proportional widths
            (should sum to 1.0).  If omitted, falls back to ``headers``.
        ``headers``
            ``[str, ...]`` — column names with equal widths.  Ignored when
            ``columns`` is present.
        ``rows``
            ``[[str, ...], ...]`` — one inner list per data row.  Strings may
            contain ``<br>`` for multi-line cells.  Alternatively, rows may be
            dicts keyed by header name.
        ``style``
            Optional overrides:

            * ``bold_first_col`` (bool) — bold the first column in every data row.
            * ``header_bg`` (RGBColor) — override the header background colour.
            * ``alt_row_bg`` (RGBColor) — override the alternate-row background.
            * ``font_name`` (str) — override the font family.
            * ``separator_color`` (RGBColor) — override the row-separator colour.

    colour : str or None
        ``"Blue"`` | ``"Orange"`` | ``"Green"`` — picks the header background
        from the hx theme.  Ignored if ``style.header_bg`` is set.
    position : dict or None
        ``{"left_cm", "top_cm", "width_cm", "height_cm"}``.  Defaults to
        full-width below the headline area.
    """
    # ── Normalise column definitions ──────────────────────────────────────
    columns = table_def.get("columns")
    if columns:
        headers = [c["header"] for c in columns]
        widths = [c.get("width") for c in columns]
    else:
        headers = table_def.get("headers", [])
        widths = None

    rows = table_def.get("rows", [])
    if not rows and not headers:
        return

    # Dict rows → list rows
    if rows and isinstance(rows[0], dict):
        rows = [[str(row.get(h, "")) for h in headers] for row in rows]

    num_cols = len(headers) if headers else (max(len(r) for r in rows) if rows else 0)
    if num_cols == 0:
        return

    if widths is None:
        widths = [1.0 / num_cols] * num_cols

    has_header = bool(headers)
    num_data_rows = len(rows)
    num_table_rows = num_data_rows + (1 if has_header else 0)
    if num_table_rows == 0:
        return

    # ── Style config ──────────────────────────────────────────────────────
    style = table_def.get("style", {})
    bold_first_col = style.get("bold_first_col", False)
    font_name = style.get("font_name", "Arial")
    separator_color = style.get("separator_color", SEPARATOR)
    alt_row_bg = style.get("alt_row_bg", LIGHT_BG)

    slide_bg_hex = style.get("slide_bg")
    if slide_bg_hex:
        h = slide_bg_hex.lstrip("#")
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(
            int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    header_bg = style.get("header_bg")
    if header_bg is None:
        colour_key = colour.lower() if colour else None
        header_bg = _HEADER_COLORS.get(colour_key, DARK_TEXT)

    # ── Position ──────────────────────────────────────────────────────────
    pos = position or DEFAULT_POSITION
    left = Cm(pos["left_cm"])
    top = Cm(pos["top_cm"])
    width = Cm(pos["width_cm"])
    total_height = Cm(pos.get("height_cm", 12.5))

    fonts = _get_font_sizes(num_data_rows)

    # ── Create table shape ────────────────────────────────────────────────
    table_shape = slide.shapes.add_table(
        num_table_rows, num_cols, left, top, width, total_height,
    )
    table = table_shape.table
    table.first_row = False
    table.horz_banding = False
    table.vert_banding = False

    # Column widths
    total_width_emu = int(Cm(pos["width_cm"]))
    for col_idx in range(num_cols):
        w = widths[col_idx] if col_idx < len(widths) and widths[col_idx] else 1.0 / num_cols
        table.columns[col_idx].width = int(total_width_emu * w)

    # ── Row heights — compact header, body rows share the rest ────────
    if has_header and num_data_rows > 0:
        header_h = int(Cm(1.0))
        body_h = (int(total_height) - header_h) // num_data_rows
        table.rows[0].height = header_h
        for r in range(1, num_table_rows):
            table.rows[r].height = body_h

    # ── Header row ────────────────────────────────────────────────────────
    row_offset = 0
    if has_header:
        for col_idx, header_text in enumerate(headers):
            cell = table.cell(0, col_idx)
            _style_cell(cell, header_text, fonts["header"], bold=True,
                        bg_color=header_bg, font_color=WHITE, font_name=font_name,
                        v_anchor=MSO_ANCHOR.MIDDLE)
            _remove_cell_borders(cell)
        row_offset = 1

    # ── Data rows ─────────────────────────────────────────────────────────
    for row_idx, row_data in enumerate(rows):
        bg = WHITE if row_idx % 2 == 0 else alt_row_bg
        for col_idx in range(num_cols):
            cell_text = row_data[col_idx] if col_idx < len(row_data) else ""
            cell = table.cell(row_idx + row_offset, col_idx)
            is_bold = bold_first_col and col_idx == 0
            _style_cell(cell, cell_text, fonts["body"], bold=is_bold,
                        bg_color=bg, font_name=font_name)
            _remove_cell_borders(cell)
            _set_thin_border(cell, sides=["lnB"], color=separator_color)
